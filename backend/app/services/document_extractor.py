import re
from dataclasses import dataclass
from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".ppt", ".pptx", ".doc", ".docx", ".md", ".markdown", ".txt"}
MAX_BLOCK_CHARS = 40_000


class DocumentExtractionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    locator: dict


def extract_document(path: Path, extension: str) -> list[ExtractedBlock]:
    extension = extension.casefold()
    extractors = {
        ".pdf": _extract_pdf,
        ".pptx": _extract_pptx,
        ".ppt": _extract_legacy_powerpoint,
        ".docx": _extract_docx,
        ".doc": _extract_legacy_word,
        ".md": _extract_markdown,
        ".markdown": _extract_markdown,
        ".txt": _extract_text,
    }
    if extension not in extractors:
        raise DocumentExtractionError("不支持这种文件格式")

    try:
        blocks = extractors[extension](path)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError("文件损坏、加密或格式无法识别") from error

    blocks = [ExtractedBlock(block.text.strip(), block.locator) for block in blocks if block.text.strip()]
    if not blocks:
        raise DocumentExtractionError("未提取到文字，文件可能是扫描件或内容为空")
    return blocks


def _split_text(text: str, locator: dict) -> list[ExtractedBlock]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= MAX_BLOCK_CHARS:
        return [ExtractedBlock(text, locator)]

    blocks = []
    start = 0
    part = 1
    while start < len(text):
        end = min(start + MAX_BLOCK_CHARS, len(text))
        if end < len(text):
            line_break = text.rfind("\n", start, end)
            if line_break > start:
                end = line_break + 1
        block_locator = {**locator, "part": part}
        blocks.append(ExtractedBlock(text[start:end].strip(), block_locator))
        start = end
        part += 1
    return blocks


def _extract_pdf(path: Path) -> list[ExtractedBlock]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise DocumentExtractionError("PDF 已加密，无法读取文字")
        except DocumentExtractionError:
            raise
        except Exception as error:
            raise DocumentExtractionError("PDF 已加密，无法读取文字") from error

    blocks = []
    for page_number, page in enumerate(reader.pages, start=1):
        blocks.extend(_split_text(page.extract_text() or "", {"page": page_number}))
    return blocks


def _shape_text(shape) -> list[str]:
    texts = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            texts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                texts.append(row_text)
    if getattr(shape, "shape_type", None) == 6:
        for child in shape.shapes:
            texts.extend(_shape_text(child))
    return texts


def _extract_pptx(path: Path) -> list[ExtractedBlock]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(value for shape in slide.shapes for value in _shape_text(shape))
        blocks.extend(_split_text(text, {"slide": slide_number}))
    return blocks


def _extract_docx(path: Path) -> list[ExtractedBlock]:
    from docx import Document

    document = Document(str(path))
    blocks = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        blocks.extend(_split_text(paragraph.text, {"paragraph": paragraph_number}))
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            blocks.extend(_split_text(text, {"table": table_number, "row": row_number}))
    return blocks


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "big5"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("文本编码无法识别")


def _extract_markdown(path: Path) -> list[ExtractedBlock]:
    lines = _decode_text(path).splitlines()
    headings = [index for index, line in enumerate(lines) if re.match(r"^#{1,6}\s+\S", line)]
    if not headings:
        return _split_text("\n".join(lines), {"start_line": 1, "end_line": len(lines)})

    starts = ([0] if headings[0] else []) + headings
    blocks = []
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(lines)
        heading = lines[start].lstrip("# ").strip() if start in headings else None
        locator = {"start_line": start + 1, "end_line": end}
        if heading:
            locator["heading"] = heading
        blocks.extend(_split_text("\n".join(lines[start:end]), locator))
    return blocks


def _extract_text(path: Path) -> list[ExtractedBlock]:
    text = _decode_text(path)
    return _split_text(text, {"start_line": 1, "end_line": len(text.splitlines())})


def _extract_legacy_word(path: Path) -> list[ExtractedBlock]:
    try:
        import pythoncom
        import win32com.client
    except ImportError as error:
        raise DocumentExtractionError("读取 .doc 需要在 Windows 安装 Microsoft Word 和 pywin32") from error

    pythoncom.CoInitialize()
    application = None
    document = None
    try:
        application = win32com.client.DispatchEx("Word.Application")
        application.Visible = False
        application.DisplayAlerts = 0
        application.AutomationSecurity = 3
        document = application.Documents.Open(
            FileName=str(path.resolve()),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
            Visible=False,
            OpenAndRepair=False,
            NoEncodingDialog=True,
        )
        text = document.Content.Text.replace("\x07", "\n").replace("\r", "\n")
        return _split_text(text, {"section": "document"})
    except Exception as error:
        raise DocumentExtractionError("无法使用 Microsoft Word 读取 .doc 文件") from error
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def _legacy_powerpoint_shape_text(shape) -> list[str]:
    texts = []
    try:
        if shape.HasTextFrame and shape.TextFrame.HasText:
            text = shape.TextFrame.TextRange.Text.strip()
            if text:
                texts.append(text)
    except Exception:
        pass
    try:
        if shape.HasTable:
            table = shape.Table
            for row_number in range(1, table.Rows.Count + 1):
                row = []
                for column_number in range(1, table.Columns.Count + 1):
                    text = table.Cell(row_number, column_number).Shape.TextFrame.TextRange.Text.strip()
                    if text:
                        row.append(text)
                if row:
                    texts.append("\t".join(row))
    except Exception:
        pass
    try:
        if shape.Type == 6:
            for index in range(1, shape.GroupItems.Count + 1):
                texts.extend(_legacy_powerpoint_shape_text(shape.GroupItems.Item(index)))
    except Exception:
        pass
    return texts


def _extract_legacy_powerpoint(path: Path) -> list[ExtractedBlock]:
    try:
        import pythoncom
        import win32com.client
    except ImportError as error:
        raise DocumentExtractionError("读取 .ppt 需要在 Windows 安装 Microsoft PowerPoint 和 pywin32") from error

    pythoncom.CoInitialize()
    application = None
    presentation = None
    try:
        application = win32com.client.DispatchEx("PowerPoint.Application")
        application.AutomationSecurity = 3
        presentation = application.Presentations.Open(str(path.resolve()), -1, 0, 0)
        blocks = []
        for slide_number in range(1, presentation.Slides.Count + 1):
            slide = presentation.Slides.Item(slide_number)
            texts = []
            for shape_number in range(1, slide.Shapes.Count + 1):
                texts.extend(_legacy_powerpoint_shape_text(slide.Shapes.Item(shape_number)))
            blocks.extend(_split_text("\n".join(texts), {"slide": slide_number}))
        return blocks
    except Exception as error:
        raise DocumentExtractionError("无法使用 Microsoft PowerPoint 读取 .ppt 文件") from error
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()
