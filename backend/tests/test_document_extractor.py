from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services import document_extractor
from app.services.document_extractor import ExtractedBlock, extract_document


def _write_text_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 18 Tf 72 720 Td (PDF course text) Tj ET")
    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


def test_extracts_pdf_pptx_docx_markdown_and_txt(tmp_path):
    pdf_path = tmp_path / "course.pdf"
    _write_text_pdf(pdf_path)

    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Transformer"
    slide.placeholders[1].text = "Self-attention content"
    presentation.save(pptx_path)

    docx_path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Word paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.save(docx_path)

    markdown_path = tmp_path / "lesson.md"
    markdown_path.write_text("# 第一章\n注意力机制\n## 第二节\n位置编码", encoding="utf-8")
    text_path = tmp_path / "plain.txt"
    text_path.write_text("纯文本课程内容", encoding="utf-8")

    pdf = extract_document(pdf_path, ".pdf")
    pptx = extract_document(pptx_path, ".pptx")
    docx = extract_document(docx_path, ".docx")
    markdown = extract_document(markdown_path, ".md")
    text = extract_document(text_path, ".txt")

    assert "PDF course text" in pdf[0].text
    assert pdf[0].locator == {"page": 1}
    assert "Transformer" in pptx[0].text and pptx[0].locator == {"slide": 1}
    assert any(block.text == "Word paragraph" for block in docx)
    assert any("left\tright" in block.text for block in docx)
    assert [block.locator["heading"] for block in markdown] == ["第一章", "第二节"]
    assert text == [ExtractedBlock("纯文本课程内容", {"start_line": 1, "end_line": 1})]


def test_legacy_doc_and_ppt_dispatch_to_office_extractors(tmp_path, monkeypatch):
    doc_path = tmp_path / "old.doc"
    ppt_path = tmp_path / "old.ppt"
    doc_path.write_bytes(b"doc")
    ppt_path.write_bytes(b"ppt")
    monkeypatch.setattr(
        document_extractor,
        "_extract_legacy_word",
        lambda path: [ExtractedBlock(f"word:{path.name}", {"section": "document"})],
    )
    monkeypatch.setattr(
        document_extractor,
        "_extract_legacy_powerpoint",
        lambda path: [ExtractedBlock(f"ppt:{path.name}", {"slide": 1})],
    )

    assert extract_document(doc_path, ".doc")[0].text == "word:old.doc"
    assert extract_document(ppt_path, ".ppt")[0].text == "ppt:old.ppt"
